"""Generate final paper supplements only from completed, audited run records."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font
from openpyxl.utils import get_column_letter
from common.result_io import write_result
from q2_archive import save_json
from q2_export import SELECTED,emergency_intervals
from q2_data import load_data

def main(root=Path('result/formal_v2')):
    root=Path(root)
    read=lambda p:json.loads(p.read_text(encoding='utf-8'))
    from q2_stress_audit import main as stress_audit
    stress_audit(root)
    summary=read(root/'summary_2025-12-31.json')
    audit=read(root/'audit.json')
    assert summary['complete'] and audit['annual_complete']
    cfg=read(root/'config.json')
    freeze=read(root/'freeze.json')
    cold=read(root/'warmup/2025-01-01.json')
    paper=root.parent/'paper'
    paper.mkdir(exist_ok=True)
    write_result(pd.DataFrame(summary['warmup']),paper/'warmup_daily.xlsx')
    write_result(pd.DataFrame([dict(date='2025-01-01',cash_cost=cold['cash_cost'],policy=cold['policy'])]),paper/'cold_start_separate.xlsx')
    ess=[]
    for system in ['star','dagger']:
        for path in sorted((root/'scenarios'/system).glob('*.json')):
            meta=read(path)
            ess.append(dict(date=path.stem,system=system,scenarios=len(meta['dates']),
                            effective_n=meta['weight_info']['effective_n']))
    df=pd.DataFrame(ess)
    write_result(df,paper/'effective_sample_size_daily.xlsx')
    write_result(df.groupby('system')[['scenarios','effective_n']].agg(['min','mean','max']).reset_index().pipe(
        lambda x:x.set_axis(['system','scenario_min','scenario_mean','scenario_max','ESS_min','ESS_mean','ESS_max'],axis=1)),
        paper/'effective_sample_size_summary.xlsx')
    pressure=[]
    stressroot=root.parent/(root.name+'_stress')
    assert read(stressroot/'isolation_check.json')['normal_checkpoint_unchanged']
    for branch in ['load_up','pv_down','joint_multiday']:
        s=read(stressroot/branch/'summary.json')
        for strategy,rows in s['rows'].items():
            for row in rows:
                pressure.append(dict(shock=branch,strategy=strategy,gamma=cfg['stress_gamma'],**row))
    write_result(pd.DataFrame(pressure),paper/'stress_daily.xlsx')
    write_result(pd.DataFrame(pressure).groupby(['shock','strategy'])[
        ['plan_cost','emergency_cost','cash_cost','emergency_kwh','cash_change_vs_normal','emergency_fee_change_vs_normal']
        ].sum().reset_index(),paper/'stress_comparison.xlsx')
    sensitivity=root.parent/(root.name+'_sensitivity')
    profiles=pd.DataFrame(read(sensitivity/'profiles.json'))
    stability=profiles[profiles.experiment=='search_budget_x3'].copy()
    stability['objective_change']=stability.objective-stability.baseline_objective
    stability['cash_change']=stability.cash_cost-stability.baseline_cash
    write_result(stability,paper/'search_stability.xlsx')
    # Table 3 across the four named dates; use actual contiguous emergency intervals.
    price=load_data().price
    for folder in [*(root/'exports').iterdir(),root.parent/'selected_days']:
        if not folder.is_dir():
            continue
        strategy=read(folder/'export_validation.json')['strategy']
        book=Workbook()
        sheet=book.active
        sheet.title='表3_四日紧急购电'
        intervals=[]
        for i,date in enumerate(SELECTED):
            with np.load(root/'runs'/strategy/(date+'.npz')) as record:
                intervals.append(emergency_intervals(record['H'],price))
            c=2*i+1
            sheet.cell(1,c,date)
            sheet.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+1)
            sheet.cell(2,c,'同日连续区间')
            sheet.cell(2,c+1,'紧急购电量 (kWh)')
            for r,(label,kwh,_) in enumerate(intervals[-1],3):
                sheet.cell(r,c,label)
                sheet.cell(r,c+1,kwh)
        sheet.freeze_panes='A3'
        sheet.sheet_view.showGridLines=False
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width=26
        for row in sheet:
            for cell in row:
                cell.font=Font(name='Microsoft YaHei',bold=cell.row<3)
                cell.alignment=Alignment(horizontal='center',vertical='center')
                if isinstance(cell.value,(float,int)):
                    cell.number_format='0.0000'
        path=folder/'four_dates_table3.xlsx'
        book.save(path)
        checked=load_workbook(path,data_only=True).active
        for i,items in enumerate(intervals):
            assert abs(sum(checked.cell(r,2*i+2).value or 0 for r in range(3,checked.max_row+1))-sum(x[1] for x in items))<1e-6
    tests=read(Path('result/verification/acceptance_tests.json'))
    assert tests['passed']
    lines=['# 第二问实际交付结果','',
           '本报告仅汇总 formal_v2 已完成的运行记录。所有金额单位为元、电量为 kWh；下表不含单列的1月预热费用。','',
           f"原始365天，每日144槽；正式334天，六套策略各48096槽。自动测试 {tests['tests_run']} 项通过，全年独立重算与导出重读通过。",'',
           f"2月1日共同起始储电量：{freeze['shared_E0']:.6f} kWh。1月验证冻结的主预测器组成：{freeze['star']}（0=F0，1=CatBoost）；F0方法为 {freeze['f0_methods']}。",'',
           '| 策略 | 计划费 | 紧急费 | 总现金费用 | 紧急电量 | 日总费CVaR95 | 日紧急费CVaR95 |',
           '|---|---:|---:|---:|---:|---:|---:|']
    for strategy,v in summary['economic'].items():
        lines.append('| '+strategy+' | '+' | '.join(f'{v[k]:.4f}' for k in
            ['plan_cost','emergency_cost','cash_cost','emergency_kwh','cash_cvar95','emergency_cvar95'])+' |')
    lines+=['',f"CatBoost达到最短历史要求后回退/失败日期数：{len(audit['training_failure_dates'])}。",'',
        '正式 result2 位于 result/result2.xlsx，单一策略由事前配置 star_M3 指定，不根据全年费用切换。每天的计划购电量保存 G；全天购电费按用户确认写计划费加紧急费；表头只在输出副本校正。','',
        '所有正常回测、负载上冲、光伏下降、多日联合压力及敏感性结果分目录保存。压力强度和窗口由配置事先确定；多日分支继承受冲击状态，正常检查点哈希未变。','',
        '搜索预算每策略每日1156次，四起点各完成一轮144分量正负扫描；多数搜索预算耗尽，结果是可行近似解。步长缩小规则已实现，但不能声称正式预算内遍历了全部步长或已收敛。','',
        f"四个固定日期的预算三倍稳定性检查共 {len(stability)} 例，目标变化范围 [{stability.objective_change.min():.6f}, {stability.objective_change.max():.6f}] 元，真实日费用变化范围 [{stability.cash_change.min():.6f}, {stability.cash_change.max():.6f}] 元。差距较小的模型比较仍受搜索精度影响。",'',
        'TV半径和终端价值敏感性仅为四个固定日期的事后诊断，没有回写正式参数或计划，也不是全年敏感性回测。核带宽、收缩系数、残差窗口和终端系数保留工程初值，未声称是历史验证选出的最优参数。','',
        '问题一回归费用、约束、导出关键值保持一致，必要公共函数已提取至common，逻辑保持不变；问题一备份和旧输出已清理。任务期间原 plot_q1.py 出现外部绘图更新，单独记录哈希差异，不覆盖该外部改动。', '',
        '复现命令、模块复用和全部存档说明见 q2/README.md。旧开发实验已清理，正式结果均来自独立完成的formal_v2。']
    (paper/'交付结果说明.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    save_json(root/'DELIVERY_STATUS.json',dict(complete=True,annual_days=334,strategy_count=6,
        tests=tests,annual_audit=True,exports_reread=True,stress_isolation=True,
        sensitivity_cases=len(profiles),official_strategy=cfg['formal_strategy'],
        limitations=['budget-limited feasible approximation','engineering defaults not tuned optima',
                     'four-date retrospective sensitivity, not annual sensitivity']))

if __name__=='__main__':
    main()
