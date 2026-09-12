"""Assemble five paper figures and unchanged, verified submission workbooks."""
from __future__ import annotations

import hashlib
import html
import json
import os
import argparse
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

_matplotlib_config = Path(__file__).resolve().parents[2] / 'output' / 'paper_matplotlib'
_matplotlib_config.mkdir(parents=True, exist_ok=True)
os.environ['MPLCONFIGDIR'] = str(_matplotlib_config)
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import numpy as np
from openpyxl import load_workbook

from microgrid.statistical_experiment import load_run
from paper_result_analysis import make_analyses

CODE = Path(__file__).resolve().parents[2]
SOURCE = CODE / 'q2' / '结果' / 'statistical_risk_v1'
OUTPUT = CODE / '论文用结果'
DATES = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
GREEN, BLUE, RED = '#16755a', '#567696', '#b94f48'


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit(template, submitted):
    """Check template coordinates and independently reconcile exported records."""
    tw = load_workbook(template, data_only=True)
    ow = load_workbook(submitted, data_only=True)
    assert tw.sheetnames == ow.sheetnames == ['计划购电量', '充放电量', '紧急购电量']
    tp, op = tw['计划购电量'], ow['计划购电量']
    assert (tp.max_row, tp.max_column) == (op.max_row, op.max_column) == (335, 147)
    assert [tp.cell(r, 1).value for r in range(2,336)] == [op.cell(r,1).value for r in range(2,336)]
    assert [tp.cell(1,c).value for c in [1,146,147]] == [op.cell(1,c).value for c in [1,146,147]]
    for name in ['充放电量','紧急购电量']:
        assert [c.value for c in tw[name][1]] == [c.value for c in ow[name][1]]
    count, cash, h_total, g_total = 0, 0., 0., 0.
    day_arrays = {}
    for r in range(2,336):
        date = op.cell(r,1).value.date().isoformat()
        with np.load(SOURCE/'evaluation'/'M3'/f'{date}.npz') as saved:
            a = {k:saved[k] for k in saved.files}
        day_arrays[date] = a
        meta = read(SOURCE/'evaluation'/'M3'/f'{date}.json')
        assert digest(SOURCE/'evaluation'/'M3'/f'{date}.npz') == meta['array_sha256']
        np.testing.assert_allclose([op.cell(r,c).value for c in range(2,146)], a['G'], rtol=0, atol=1e-7)
        assert abs(op.cell(r,146).value-a['G'].sum()) < 1e-7
        assert abs(op.cell(r,147).value-meta['cash_cost']) < 1e-7
        cash += op.cell(r,147).value
        g_total += op.cell(r,146).value
        count += 146
        start = 2+(r-2)*6
        st = ow['充放电量']
        assert st.cell(start,1).value.date().isoformat() == date
        for k in range(6):
            assert st.cell(start+k,2).value == f'{k*4}:00-{(k+1)*4}:00'
            for c,key in [(3,'charge'),(4,'discharge')]:
                assert abs(st.cell(start+k,c).value-a[key][24*k:24*(k+1)].sum()) < 1e-7
                count += 1
        assert abs(st.cell(start,6).value-a['E'][0]) < 1e-7
        assert abs(st.cell(start+1,6).value-a['E'][-1]) < 1e-7
        count += 2
    assert ow['充放电量'].max_row == 1+334*6
    covered = {d:np.zeros(144,dtype=bool) for d in day_arrays}
    dates_present = set()
    current = None
    for values in ow['紧急购电量'].iter_rows(min_row=2,values_only=True):
        stamp, interval, amount = values
        if stamp is not None:
            current = stamp.date().isoformat()
            dates_present.add(current)
        actual = day_arrays[current]['H']
        if interval == '无':
            assert actual.max() <= 1e-6 and amount == 0
        else:
            def slot(s):
                hour, minute = map(int,s.split(':'))
                assert minute % 10 == 0
                return hour*6+minute//10
            start,end = map(slot,interval.split('-'))
            assert 0 <= start < end <= 144
            assert not covered[current][start:end].any()
            assert (actual[start:end] > 1e-6).all()
            assert abs(amount-actual[start:end].sum()) < 1e-7
            covered[current][start:end] = True
        h_total += amount
        count += 1
    assert dates_present == set(day_arrays)
    for d,a in day_arrays.items():
        np.testing.assert_array_equal(covered[d],a['H'] > 1e-6)
    summary = read(SOURCE/'evaluation'/'M3'/'summary.json')
    for actual,key in [(cash,'cash_cost'),(h_total,'H'),(g_total,'G')]:
        assert abs(actual-summary[key]) < 1e-6
    differences = [
        {'cell':tp.cell(1,c).coordinate,'template':tp.cell(1,c).value,'output':op.cell(1,c).value}
        for c in range(2,146) if tp.cell(1,c).value != op.cell(1,c).value
    ]
    result = {
        'created':datetime.now().isoformat(), 'template':str(template),
        'template_sha256':digest(template), 'submitted_sha256':digest(submitted),
        'sheet_names_and_order_match':True, 'all_334_dates_match':True,
        'numeric_values_reconciled':count, 'cash_cost':cash,'planned_kwh':g_total,'emergency_kwh':h_total,
        'storage_rows':ow['充放电量'].max_row,'emergency_rows':ow['紧急购电量'].max_row,
        'time_header_exact_match':False,'time_header_differences':differences,
        'interpretation':'工作表、列结构、日期及数值映射通过核验；144个计划时段表头不逐字相同。输出采用当日00:00—24:00，不采用示例从00:10起延伸到次日00:10的标签。充放电和紧急记录按334天扩展。',
    }
    tw.close()
    ow.close()
    return result


def build():
    cfg, protocol = load_run(SOURCE)
    assert read(SOURCE/'selection_frozen.json')['selected'] == 'M3'
    template = CODE.parent/'附件5'/'result2.xlsx'
    verification = audit(template, SOURCE/'submission'/'M3'/'result2.xlsx')
    OUTPUT.mkdir(exist_ok=True)
    figures = OUTPUT/'正文图'
    figures.mkdir(exist_ok=True)
    vector = SOURCE/'paper_package_support'/'矢量图'
    vector.mkdir(parents=True,exist_ok=True)
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],
                        'axes.unicode_minus':False,'font.size':11,'svg.fonttype':'none'})
    entries = []

    def save(fig, stem, section, purpose, caption):
        fig.savefig(figures/f'{stem}.png',dpi=300,bbox_inches='tight')
        fig.savefig(vector/f'{stem}.svg',bbox_inches='tight')
        plt.close(fig)
        entries.append(dict(name=stem,section=section,purpose=purpose,caption=caption))

    models = ['M0','M1','M2','M3']
    summaries = {m:read(SOURCE/'evaluation'/m/'summary.json') for m in models}
    fig, axes = plt.subplots(1,3,figsize=(13,4.4),layout='constrained')
    for ax,key,label,scale in zip(axes,['cash_cost','H','emergency_cvar'],
            ['总现金费用 / 万元','紧急购电量 / 万 kWh',r'日紧急费用 $\mathrm{CVaR}_{0.9}$ / 元'],[1e4,1e4,1]):
        values = [summaries[m][key]/scale for m in models]
        bars=ax.barh(models,values,color=[GREEN if m=='M3' else BLUE for m in models])
        for bar,value in zip(bars,values):
            label_x=max(values)*1.065 if key=='cash_cost' else value+max(values)*.015
            ax.text(label_x,bar.get_y()+bar.get_height()/2,f'{value:,.2f}',va='center',fontsize=10)
        ax.set(xlabel=label,xlim=(0,max(values)*1.28))
        ax.invert_yaxis();ax.grid(axis='x',alpha=.18)
        if key=='cash_cost':
            ax.axvline(summaries['M2'][key]*1.03/scale,ls='--',color='#bc6b25',lw=1.2,
                       label='M2 费用 × 1.03')
            ax.legend(loc='lower right',fontsize=8,frameon=False)
    fig.suptitle('M0—M3 递进消融：2025 年 2—12 月（334 天）')
    save(fig,'图1_M0至M3消融','模型有效性与消融实验','解释误差场景、条件权重和 TV 鲁棒的增量作用。',
         '各模型共用统计预测、物理控制和搜索预算。绿色为一月选定的 M3；费用虚线仅为后续评价的 3% 核验线，不表示逐日费用约束。CVaR 为真实日紧急费用的经验尾部统计量。')

    days={}
    for date in DATES:
        days[date]={}
        for model in ['M2','M3']:
            with np.load(SOURCE/'evaluation'/model/f'{date}.npz') as a:
                days[date][model]={k:a[k] for k in a.files}
    fig,axes=plt.subplots(2,2,figsize=(12,7.5),layout='constrained')
    for ax,date in zip(axes.flat,DATES):
        a=days[date]['M3']
        ax.stairs(a['load']-a['pv'],np.arange(145)/6,baseline=None,color='#777777',lw=1,label='实际净负荷 L−V')
        ax.stairs(a['G'],np.arange(145)/6,baseline=None,color=GREEN,lw=1.6,label='M3 日前计划 G')
        ax.stairs(days[date]['M2']['G'],np.arange(145)/6,baseline=None,color=BLUE,lw=1,ls='--',label='M2 日前计划 G')
        ax.stairs(a['H'],np.arange(145)/6,color=RED,fill=True,alpha=.35,label='M3 紧急购电 H')
        ax.axhline(0,color='#aaaaaa',lw=.5)
        ax.set(title=date,xlim=(0,24),xticks=range(0,25,4),xlabel='时刻 / h',ylabel='每十分钟电量 / kWh')
        ax.grid(alpha=.18)
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc='outside lower center',ncol=4,frameon=False,fontsize=10)
    fig.suptitle('四个指定日：日前购电计划与实际供需')
    save(fig,'图2_四个指定日购电计划','购电结果与决策解释','把原来四张分散的指定日购电图合成一张。',
         '四个子图对应题设指定日期，所有曲线使用同一电量单位，未使用双纵轴。实际净负荷可为负，表示光伏超过负载。实际曲线用于事后解释；日初计划仅依赖过去信息。红色为 M3 紧急购电量。')

    fig,axes=plt.subplots(2,2,figsize=(12,7.5),layout='constrained')
    for ax,date in zip(axes.flat,DATES):
        a=days[date]['M3']
        ax.plot(np.arange(145)/6,a['E']/1000,color=GREEN,lw=1.7,label='M3 实际库存')
        ax.plot(np.arange(145)/6,days[date]['M2']['E']/1000,color=BLUE,lw=1.1,label='M2 实际库存')
        ax.plot(np.arange(145)/6,a['reference']/1000,color='#b98540',lw=1.1,ls='--',label='M3 参考库存')
        ax.axhline(cfg.e_min/1000,color='#888888',ls=':',lw=1)
        ax.axhline(cfg.e_max/1000,color='#888888',ls=':',lw=1)
        ax.set(title=date,xlim=(0,24),ylim=(0,12),xticks=range(0,25,4),xlabel='时刻 / h',ylabel='储电量 / 千 kWh')
        ax.grid(alpha=.18)
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc='outside lower center',ncol=3,frameon=False)
    fig.suptitle('四个指定日：储能库存与参考库存（虚点线为物理上下界）')
    save(fig,'图3_四个指定日储能响应','储能调节机理','配合图2解释放电和留电时机。',
         '参考库存是控制器的放电保留阈值，不要求实际库存始终高于参考线。各模型从共同评价初始状态出发，此后沿用各自真实日末库存，因此指定日的日初库存可以不同。库存不能按时段电量乘 6 换算为功率。')

    rows=[read(p) for p in sorted((SOURCE/'evaluation'/'M3').glob('????-??-??.json'))]
    dates=[datetime.fromisoformat(r['date']) for r in rows]
    fig,ax=plt.subplots(figsize=(12,4.6),layout='constrained')
    ax.plot(dates,[r['forecast_risk']['emergency_cvar'] for r in rows],color='#b98540',lw=1,label=r'事前 $\mathrm{CVaR}_{0.9}$')
    ax.plot(dates,[r['forecast_risk']['emergency_var'] for r in rows],color=BLUE,lw=1,label=r'事前 $\mathrm{VaR}_{0.9}$')
    ax.plot(dates,[r['emergency_cost'] for r in rows],color=RED,lw=.9,label='真实日紧急费用')
    ax.set(xlim=(dates[0],dates[-1]),ylim=(0,None),ylabel='日紧急费用 / 元',xlabel='2025 年日期')
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%m月'))
    ax.grid(alpha=.18)
    ax.legend(loc='upper right',frameon=False,ncol=3,fontsize=10)
    ax.set_title('M3 风险评估：实际费用超过事前 VaR 的比例为 12.57%')
    save(fig,'图4_事前风险与实际费用','风险评估与校准','展示风险提示与实测损失的对应关系。',
         '事前 VaR/CVaR 仅使用当时可用的历史误差场景。真实费用超过事前 VaR 的比例为 42/334=12.57%，名义参考为 10%。这是一项校准诊断，CVaR 不是每日费用上界。')

    stress=read(SOURCE/'stress_summary.json')['models']
    shocks=protocol['heldout_shocks']
    titles=['负载 +20%','光伏 −50%','17—21 时负载 +30%','负载 +15%、光伏 −40%','三日：负载 +15%、光伏 −40%','低库存三日：负载 +10%、光伏 −30%']
    fig,axes=plt.subplots(2,3,figsize=(13,7.8),layout='constrained')
    for ax,shock,title in zip(axes.flat,shocks,titles):
        for i,model in enumerate(['M2','M3']):
            lookup={r['anchor']:r for r in stress[model]['results'] if r['shock']['name']==shock['name']}
            ax.bar(np.arange(4)+(i-.5)*.36,[lookup[d]['emergency_kwh']/1000 for d in DATES],width=.34,
                   color=BLUE if model=='M2' else GREEN,label=model)
        ax.set(title=title,xticks=range(4),xticklabels=[d[5:] for d in DATES],
               xlabel='冲击起始日',ylabel='案例紧急电量 / 千 kWh')
        ax.title.set_fontsize(10)
        ax.grid(axis='y',alpha=.18)
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc='outside lower center',ncol=2,frameon=False)
    fig.suptitle('留出压力测试：四季起点、六类冲击')
    save(fig,'图5_六类压力测试','压力抗性与局限','比较相同冲击下的应急依赖，保留退化案例。',
         '各面板纵轴独立。联合三日冲击为负载 +15%、光伏 −40%；低库存三日为负载 +10%、光伏 −30%。两类均包含两天恢复期。低库存初值为 2160 kWh，其余为 6000 kWh。首日执行冲击前同一锁定计划。人工冲击没有发生概率，案例合计不能解释为年期望损失；当前允许无限紧急补购。')

    for name in ['result2.xlsx','论文结果表.xlsx']:
        shutil.copy2(SOURCE/'submission'/'M3'/name,OUTPUT/name)
        assert digest(OUTPUT/name) == digest(SOURCE/'submission'/'M3'/name)
    verification_file=SOURCE/'audit'/'提交模板对齐核验.json'
    verification_file.write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf-8')
    archive=OUTPUT/'重绘资料.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for e in entries:
            z.write(vector/(e['name']+'.svg'),'矢量图/'+e['name']+'.svg')
        for p in sorted((SOURCE/'figures'/'绘图数据').glob('*.csv')):
            z.write(p,'绘图数据/'+p.name)
        for name in ['evaluation_comparison.csv','heldout_stress_cases.csv']:
            z.write(SOURCE/'report'/name,'绘图数据/'+name)
        for name in ['论文模型分析写作稿.md','论文补充分析.json','模型与风险结果.md']:
            z.write(SOURCE/'report'/name,'分析与核验/'+name)
        z.write(verification_file,'分析与核验/'+verification_file.name)
        z.writestr('使用说明.txt','PNG 正文图位于压缩包外。矢量图可用于排版，CSV 可用于重绘。SVG 使用本机微软雅黑字体。十分钟流量为 kWh，乘6可换算平均kW；库存不乘6。分析文稿原版位于 q2/结果/statistical_risk_v1/report，相对链接请从原目录打开。')
    nav=make_navigation(entries,verification)
    (OUTPUT/'00_从这里开始.html').write_text(nav,encoding='utf-8')
    manifest={'created':datetime.now().isoformat(),'model':'M3','source':str(SOURCE),
              'optimization_rerun':False,'workbooks_unchanged':True,'main_figures':entries,
              'files':{str(p.relative_to(OUTPUT)):digest(p) for p in OUTPUT.rglob('*') if p.is_file()}}
    (SOURCE/'audit'/'论文用结果清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return {'output':str(OUTPUT),'main_figures':len(entries),'template_audit':{
        k:verification[k] for k in ['all_334_dates_match','numeric_values_reconciled','time_header_exact_match','cash_cost','emergency_kwh']}}


def make_navigation(entries,check):
    cards=[]
    analyses=make_analyses(SOURCE)
    modeling_example=(Path(__file__).resolve().parent/'论文写作'/'问题二建模示例.html').read_text(encoding='utf-8')
    for e in entries:
        name=html.escape(e['name'])
        cards.append(f'<article id="{name}"><h2>{name.replace("_"," · ")}</h2><p class="place">放在：{html.escape(e["section"])}</p>'
                     f'<p>{html.escape(e["purpose"])}</p><a href="正文图/{name}.png"><img src="正文图/{name}.png" alt="{name}"></a>'
                     f'<p class="caption"><b>建议图注：</b>{html.escape(e["caption"])}</p>'
                     f'<div class="figure-analysis">{analyses[e["name"].split("_",1)[0]]}</div></article>')
    mapping=''.join(f'<tr><td>{i}</td><td><a href="#{html.escape(e["name"])}">{html.escape(e["name"].split("_",1)[1])}</a></td><td>{html.escape(e["section"])}</td></tr>' for i,e in enumerate(entries,1))
    return '''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>问题二论文与提交结果</title><style>
body{font:16px/1.8 "Microsoft YaHei",sans-serif;color:#213029;background:#f4f6f4;margin:0}
main{max-width:1080px;margin:36px auto;padding:0 24px 60px}h1{font-size:30px}h2{font-size:22px}
a{color:#126349}article,section{background:white;border:1px solid #e1e7e2;border-radius:10px;padding:24px;margin:24px 0}
img{width:100%;height:auto;margin-top:12px}table{border-collapse:collapse;width:100%}th,td{text-align:left;border-bottom:1px solid #e2e7e2;padding:10px}
.place{color:#557161}.caption{background:#f4f7f4;padding:14px;font-size:14px}.note{border-left:4px solid #b17a35;padding-left:18px}
code{background:#f0f3f0;padding:2px 6px}nav a{display:inline-block;margin:0 20px 8px 0}small{color:#526159}
.analysis-nav{position:sticky;top:0;background:#f4f6f4f5;backdrop-filter:blur(8px);z-index:2;padding:12px 14px 4px;border-bottom:1px solid #d7e0d8}
.analysis-nav a{font-size:15px;margin-right:18px}article{scroll-margin-top:80px}h3{font-size:21px;margin-top:30px;color:#145c43}h4{font-size:17px;margin-bottom:8px}
.figure-analysis{border-top:1px solid #e0e7e1;margin-top:24px}.table-scroll{overflow-x:auto;margin:20px 0}
.figure-analysis table{font-size:14px;line-height:1.6}.figure-analysis th{background:#f1f5f1}.figure-analysis td{vertical-align:top}
caption{text-align:left;font-size:13px;color:#526159;padding-bottom:8px}details{border:1px solid #dce5dd;border-radius:6px;padding:14px 16px;margin:20px 0}
summary{cursor:pointer;font-weight:600;color:#245d46}.paper-draft{background:#f4f7f3;border-left:4px solid #16755a;padding:8px 18px 16px;margin:24px 0}
.paper-draft h4{margin-top:8px}blockquote{margin:0;line-height:1.9}p{overflow-wrap:anywhere}
.manuscript-body{max-width:870px;margin:0 auto;font:18px/1.95 Cambria,"Times New Roman","SimSun",serif;color:#202b24}
.manuscript-body h3{font-family:"Microsoft YaHei",sans-serif;line-height:1.6;font-size:20px;scroll-margin-top:90px;margin-top:34px}
.manuscript-body p{margin:16px 0;text-align:justify}.manuscript-body table{font:15px/1.8 "Microsoft YaHei",sans-serif}
.editor-note{font:14px/1.8 "Microsoft YaHei",sans-serif;color:#536557;background:#f5f7f3;padding:12px 16px}
.manuscript-toc{padding:12px 0;border-bottom:1px solid #e1e7e2;margin-bottom:24px}.manuscript-toc a{font-size:14px}
.equation{font-family:"Cambria Math",Cambria,"Times New Roman",serif;font-size:18px;background:#fafcf9;border-left:2px solid #7b9d82;padding:14px 12px;margin:22px 0;overflow-x:auto}
.equation>span{display:block;text-align:center;min-width:max-content;white-space:nowrap;line-height:2.1}
.equation .eq-number{text-align:right;min-width:0;white-space:normal;font-size:13px;color:#5c6e60;line-height:1.4;margin-top:4px}
.equation sub,.equation sup{font-size:.75em}.equation-lines>span{padding:2px 0}
@media(max-width:650px){main{padding:0 12px 40px;margin-top:20px}article,section{padding:16px}h1{font-size:25px}.analysis-nav a{font-size:13px;margin-right:12px}.figure-analysis table{min-width:660px}}
@media print{.analysis-nav{position:static}article,section{border:0}details{break-inside:avoid}img{max-height:75vh;object-fit:contain}}
</style><main><h1>问题二论文与提交结果</h1>
<p>本次模型：<b>M3</b>。2025 年 2—12 月连续评价，共 334 天。这个目录只放两个结果表、五张正文图和一个重绘资料包。</p>
<nav><a href="result2.xlsx">打开提交结果表</a><a href="论文结果表.xlsx">打开论文表格</a><a href="重绘资料.zip">下载重绘资料</a></nav>
<nav class="analysis-nav" aria-label="结果分析导航"><a href="#q2-model-example">建模正文示例</a><a href="#图1_M0至M3消融">消融实验</a><a href="#图2_四个指定日购电计划">购电计划</a><a href="#图3_四个指定日储能响应">储能响应</a><a href="#图4_事前风险与实际费用">风险校准</a><a href="#图5_六类压力测试">压力测试</a></nav>
<p>每张图下方均包含数值对照、结果解释和可用于论文的表述。消融部分回答“为什么选 M3”，压力部分给出六类冲击汇总及全部 24 个配对案例；展开补充表格可查看细节。</p>
'''+modeling_example+'''
<section><h2>先区分两个表格</h2><table><tr><th>文件</th><th>用途</th></tr>
<tr><td>result2.xlsx</td><td>按附件 5 的三个工作表输出提交数据：计划购电量、充放电量、紧急购电量。</td></tr>
<tr><td>论文结果表.xlsx</td><td>供正文取表：评价期汇总及四个指定日的购电、充放电和紧急购电结果。它不是附件 5 提交模板。</td></tr></table>
<p>两份 Excel 都是已核验原输出的原样复制。本次整理和合图直接读取保存结果，没有重新计算购电计划。</p></section>
<section><h2>与示例表的对齐情况</h2>
<table><tr><th>检查项</th><th>结果</th></tr>
<tr><td>工作表名称与顺序</td><td>相同：计划购电量、充放电量、紧急购电量。</td></tr>
<tr><td>计划购电量</td><td>335 行 × 147 列；334 个日期全部对应；144 个十分钟计划量，末两列为全天计划购电量及全天总现金费用。</td></tr>
<tr><td>充放电量</td><td>相同六列；按每天六个四小时段填写，共 2005 行（含表头），并记录 00:00 与 24:00 库存。</td></tr>
<tr><td>紧急购电量</td><td>相同三列；按实际连续应急区间展开，共 466 行（含表头）；无应急日期记“无”和 0。</td></tr>
<tr><td>数值对应</td><td>逐时购电、分段充放电、首末库存与应急区间均与保存数组核对通过。</td></tr></table>
<p class="note"><b>时间表头并非逐字相同。</b>示例 B1 为“0:10–0:20”，EO1 为“0:00–0:10+1”；本次输出 B1 为“0:00–0:10”，EO1 为“23:50–24:00”。本次按照当日 00:00—24:00 的完整一天映射数值，未把计划整体移位。若要求逐字保留示例的时间标签，当前版不满足这一点；这里的核验不等于确认提交系统接受了该标签修正。原始示例文件未改动。</p>
<p>示例后两个工作表只有样例行；正式结果按完整日期与实际事件扩展，所以行数不同是正常的。全天购电量是日前计划 G 的合计；全天购电费包含计划费及紧急补购费，不是仅把 G 乘电价。</p></section>
<section><h2>正文只使用这五张图</h2><table><tr><th>顺序</th><th>图名</th><th>放置位置</th></tr>'''+mapping+'''</table>
<p>图号可按论文顺序调整。四个指定日已分别合并成一张购电图和一张库存图，不必再插入八张单日图。若篇幅紧张，图3可移到附录，正文保留相应储能结果表。</p>
<p>累计费用、全年供需总览、M2 单独风险图及 R/S 扩展图不需要全部放正文。主文用下方五张图建立“消融—决策—风险—压力”的证据顺序。</p></section>'''+''.join(cards)+'''
<section><h2>取数与重绘</h2><p>PNG 均为 300 dpi；同名 SVG、M2/M3 每十分钟 CSV、风险 CSV 和汇总对照表集中在 <a href="重绘资料.zip">重绘资料.zip</a>。M2、M3 每个十分钟明细均为 48096 行。</p>
<p>论文只需保留关键指标：总现金费用、紧急电量、日紧急费用 CVaR、最大紧急功率、光伏消纳率以及压力对照。α、每个历史场景权重、逐次优化日志等属于复核细节，不必逐项堆进正文。</p>
<p>本次 M3 相对 M2：费用 +0.64%，紧急电量 −16.99%，紧急费用 CVaR −20.42%，最大紧急功率 +12.41%。最终 M3 没有启用 CVaR/压力硬约束；风险改善不代表每项指标都改善。</p></section>
<small>原始运行结果仍保存在 q2/结果/statistical_risk_v1；详细模板核验在该目录 audit/提交模板对齐核验.json。此页面是论文取图与提交文件入口。</small></main></html>'''


def refresh_html():
    """Refresh discussion only; preserve plots, data and spreadsheet bytes."""
    manifest_path=SOURCE/'audit'/'论文用结果清单.json'
    manifest=read(manifest_path)
    check=read(SOURCE/'audit'/'提交模板对齐核验.json')
    navigation=make_navigation(manifest['main_figures'],check)
    display_navigation=navigation
    archive=OUTPUT/'正文图.zip'
    if archive.is_file():
        original_archive_hash=digest(archive)
        temporary=OUTPUT/'.正文图更新.tmp'
        with zipfile.ZipFile(archive) as old:
            assert '00_从这里开始.html' in old.namelist(), 'The archive has no matching HTML page'
            with zipfile.ZipFile(temporary,'w') as updated:
                for entry in old.infolist():
                    data=navigation.encode('utf-8') if entry.filename=='00_从这里开始.html' else old.read(entry)
                    updated.writestr(entry,data)
        assert digest(archive)==original_archive_hash, 'Archive changed during update'
        temporary.replace(archive)
        manifest['html_archive_updated']=True
        if not (OUTPUT/'重绘资料.zip').is_file():
            with zipfile.ZipFile(archive) as packed:
                assert '重绘资料.zip' in packed.namelist()
            display_navigation=navigation.replace('href="重绘资料.zip"','href="正文图.zip"')
            display_navigation=display_navigation.replace('下载重绘资料','打开图表资料包')
            display_navigation=display_navigation.replace(
                '<a href="正文图.zip">重绘资料.zip</a>','<a href="正文图.zip">正文图.zip 内的重绘资料.zip</a>')
            display_navigation=display_navigation.replace('和一个重绘资料包','和一个图表资料包')
    (OUTPUT/'00_从这里开始.html').write_text(display_navigation,encoding='utf-8')
    manifest['html_analysis_updated']=datetime.now().isoformat()
    example_path=Path(__file__).resolve().parent/'论文写作'/'问题二建模示例.html'
    manifest['modeling_example_sha256']=digest(example_path)
    manifest['analysis_sources']={
        str(path.relative_to(SOURCE)):digest(path)
        for path in [SOURCE/'stress_summary.json',SOURCE/'report'/'论文补充分析.json',
                     SOURCE/'selection_frozen.json']
    }
    prediction_audit=SOURCE/'audit'/'表6与本次预测误差核对.json'
    if prediction_audit.exists():
        manifest['analysis_sources'][str(prediction_audit.relative_to(SOURCE))]=digest(prediction_audit)
    manifest['files']={str(p.relative_to(OUTPUT)):digest(p) for p in OUTPUT.rglob('*') if p.is_file()}
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return {'output':str(OUTPUT/'00_从这里开始.html'),'figure_analyses':5,'optimization_rerun':False}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--html-only',action='store_true',help='Refresh written analyses without rebuilding plots or workbooks')
    args=parser.parse_args()
    print(json.dumps(refresh_html() if args.html_only else build(),ensure_ascii=False,indent=2))
