from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from q1_config import (
    PERIODS_PER_HOUR,
    RESULT_DIR,
    RESULT_TEMPLATE,
    interval_labels,
)
from q1_result_io import read_result


def build_tables(schedule):
    """把调度结果汇总为题目要求的两张表。"""
    labels = interval_labels()
    if schedule.time.tolist() != labels:
        raise ValueError("调度应按时间顺序包含全天的连续时段。")
    s = schedule.set_index("time")
    selected = ["10:00-10:10", "12:00-12:10", "14:00-14:10",
                "16:00-16:10", "18:00-18:10", "20:00-20:10"]
    table1 = [["时间段", "购电量", "时间段", "购电量", "时间段", "购电量"]]
    for group in (selected[:3], selected[3:]):
        table1.append([item for label in group for item in (label, float(s.loc[label, "grid_purchase"]))])
    daily_grid = float(s.grid_purchase.sum())
    daily_cost = float((s.price * s.grid_purchase).sum())
    table1.append(["全天购电量", None, daily_grid, "全天购电费", None, daily_cost])
    blocks = {}
    for hour in range(0, 24, 4):
        part = s.iloc[hour * PERIODS_PER_HOUR:(hour + 4) * PERIODS_PER_HOUR]
        blocks[f"{hour}:00-{hour + 4}:00"] = [
            float(part.battery_charge.sum()), float(part.battery_discharge.sum())
        ]
    table2 = [["时间段", "充电量", "放电量", "时间段", "充电量", "放电量"]]
    keys = list(blocks)
    for i in range(0, 6, 2):
        table2.append([keys[i], *blocks[keys[i]], keys[i + 1], *blocks[keys[i + 1]]])
    first, last = float(s.loc[labels[0], "energy_start"]), float(s.loc[labels[-1], "energy_end"])
    table2.append(["0:00 储电量", None, first, "24:00 储电量", None, last])
    return s, table1, table2, blocks, (first, last)


# 用题目表格的六列布局输出，保留完整精度，仅显示四位小数
def write_paper_table(rows, number, title, units, output_dir=RESULT_DIR):
    book = Workbook()
    sheet = book.active
    sheet.title = f"表{number}"
    sheet.append([title])
    sheet.merge_cells("A1:F1")
    sheet.append([units])
    sheet.merge_cells("A2:F2")
    for row in rows:
        sheet.append(row)
    last = sheet.max_row
    sheet.merge_cells(start_row=last, start_column=1, end_row=last, end_column=2)
    sheet.merge_cells(start_row=last, start_column=4, end_row=last, end_column=5)
    line = Side(style="thin", color="222222")
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11, bold=cell.row in (1, 3))
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"
            if cell.row == 3:
                cell.border = Border(top=line, bottom=line)
            elif cell.row == last:
                cell.border = Border(bottom=line)
        sheet.row_dimensions[row[0].row].height = 28
    for i in range(1, 7):
        sheet.column_dimensions[get_column_letter(i)].width = 21
    sheet.print_options.horizontalCentered = True
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:F{last}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    book.save(output_dir / f"q1_table{number}.xlsx")
    print(f"\n表{number}：{title}\n{units}")
    for row in rows:
        print("\t".join(str(Decimal(format(item, ".15g")).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))
                        if isinstance(item, (int, float)) else str(item or "") for item in row))


# 保留官方模板样式，在副本中校正整体晚10分钟的标签后按标签写入
def fill_template(s, blocks, states, output_dir=RESULT_DIR):
    book = load_workbook(RESULT_TEMPLATE)
    purchase = book["计划购电量"]
    if purchase.max_row != 145:
        raise ValueError("官方购电表应包含144条记录。")
    for row, label in enumerate(s.index, start=2):
        purchase.cell(row, 1, label)
        purchase.cell(row, 2, float(s.loc[label, "grid_purchase"]))
    battery = book["充放电量"]
    for row in range(2, 8):
        label = battery.cell(row, 1).value
        charge, discharge = blocks[label]
        battery.cell(row, 2, charge)
        battery.cell(row, 3, discharge)
    for row in range(2, battery.max_row + 1):
        moment = battery.cell(row, 4).value
        if moment in ("0:00", "24:00"):
            battery.cell(row, 5, states[0 if moment == "0:00" else 1])
    filled = output_dir / "result1.xlsx"
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0000"
    book.save(filled)
    print(f"已生成：{filled}（已校正模板中的时间标签）")


def export_tables(schedule, output_dir=RESULT_DIR, prefix=""):
    output_dir.mkdir(parents=True, exist_ok=True)
    s, table1, table2, blocks, states = build_tables(schedule)
    write_paper_table(table1, 1, "微网在指定时间段的购电量及全天的购电量和购电费", prefix + "购电量单位：kWh；购电费单位：元", output_dir)
    write_paper_table(table2, 2, "储能设备在指定时间段的充放电量及0:00和24:00的储电量", prefix + "电量单位：kWh", output_dir)
    fill_template(s, blocks, states, output_dir)


def main():
    export_tables(read_result(RESULT_DIR / "q1_schedule.xlsx"))


if __name__ == "__main__":
    main()
