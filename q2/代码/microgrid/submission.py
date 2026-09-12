"""A/B 共用的提交数据构建、Excel 导出与逐值核验。"""

from datetime import date, timedelta
from pathlib import Path
import argparse
import json
import os
import subprocess
import numpy as np
from .config import CODE_DIR, ATTACHMENTS
from .reporting import daily_results, METRICS
from .metrics import totals
from .storage import read_arrays, read_json, write_json, file_hash, evaluation_choice

DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
MODEL_NAMES = {"search_only": "A", "joint_reserve": "B", "adaptive_joint": "C"}


def clock(minutes):
    return f"{minutes // 60}:{minutes % 60:02d}"


def intervals(values):
    """合并连续应急时段，保留原始电量精度。"""
    rows = []
    t = 0
    while t < 144:
        if values[t] <= 1e-6:
            t += 1
            continue
        start = t
        while t < 144 and values[t] > 1e-6:
            t += 1
        rows.append([f"{clock(start*10)}-{clock(t*10)}", float(values[start:t].sum())])
    return rows


def make_payload(source, output, variant=None):
    source, output = Path(source).resolve(), Path(output).resolve()
    variant, directory, daily = daily_results(source, variant)
    model = MODEL_NAMES.get(variant, variant)
    selected = evaluation_choice(source)["selected"]
    definition = "固定α=1，优化购电计划" if model == "A" else "联合优化购电计划与α"
    protocol = read_json(source/"protocol_frozen.json") if (source/"protocol_frozen.json").exists() else {}
    if protocol.get("schema") == "statistical-risk-v1":
        from .models import get_model
        spec = get_model(variant)
        definition = f"统计预测；{spec.scenario_mode}场景；TV {'开' if spec.robust else '关'}；风险限额{'开' if spec.risk_limit else '关'}；压力保护{'开' if spec.stress_guard else '关'}"
        completion = read_json(directory/"complete.json")
        if not completion["complete"] or completion["days"] != 334:
            raise ValueError("Only a complete 334-day run may be submitted")
    summary = totals(daily)
    purchase = [
        ["日期\\时间"]
        + [f"{clock(t*10)}-{clock((t+1)*10)}" for t in range(144)]
        + ["全天购电量", "全天购电费"]
    ]
    storage = [["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]]
    emergency = [["日期", "购电时间段", "购电量"]]
    chosen = {}
    previous = daily[0]["E0"]
    for row in daily:
        d = row["date"]
        a = read_arrays(directory / (d + ".npz"))
        if "array_sha256" in row and file_hash(directory/(d+".npz")) != row["array_sha256"]:
            raise ValueError(f"计算结果已改变：{d}")
        if abs(a["E"][0] - previous) > 1e-6:
            raise ValueError(f"储能跨日状态不连续：{d}")
        previous = a["E"][-1]
        serial = (date.fromisoformat(d) - date(1899, 12, 30)).days
        purchase.append(
            [serial, *a["G"].tolist(), float(a["G"].sum()), row["cash_cost"]]
        )
        blocks = []
        for k in range(6):
            blocks.append(
                [
                    serial if k == 0 else None,
                    f"{k*4}:00-{(k+1)*4}:00",
                    float(a["charge"][k * 24 : (k + 1) * 24].sum()),
                    float(a["discharge"][k * 24 : (k + 1) * 24].sum()),
                    "0:00" if k == 0 else "24:00" if k == 1 else None,
                    (
                        float(a["E"][0])
                        if k == 0
                        else float(a["E"][-1]) if k == 1 else None
                    ),
                ]
            )
        storage.extend(blocks)
        events = intervals(a["H"])
        emergency.extend(
            [[serial if k == 0 else None, *event] for k, event in enumerate(events)]
            if events
            else [[serial, "无", 0.0]]
        )
        if d in DATES:
            chosen[d] = (a, row, blocks, events)
    table1, table2 = [], []
    for d in DATES:
        a, row, blocks, events = chosen[d]
        table1.extend(
            [
                [d, None, None, None, None, None],
                [
                    "时间段",
                    "购电量 / kWh",
                    "时间段",
                    "购电量 / kWh",
                    "时间段",
                    "购电量 / kWh",
                ],
            ]
        )
        for hours in ((10, 12, 14), (16, 18, 20)):
            cells = []
            for h in hours:
                cells.extend([f"{h}:00-{h}:10", float(a["G"][h * 6])])
            table1.append(cells)
        table1.extend(
            [
                [
                    "全天计划购电量",
                    row["G"],
                    "全天总购电费 / 元",
                    row["cash_cost"],
                    None,
                    None,
                ],
                [None] * 6,
            ]
        )
        table2.extend(
            [
                [d, None, None, None, None, None],
                [
                    "时间段",
                    "充电量 / kWh",
                    "放电量 / kWh",
                    "时间段",
                    "充电量 / kWh",
                    "放电量 / kWh",
                ],
            ]
        )
        for k in (0, 2, 4):
            left, right = blocks[k : k + 2]
            table2.append([*left[1:4], *right[1:4]])
        table2.extend(
            [
                [
                    "0:00储电量",
                    float(a["E"][0]),
                    None,
                    "24:00储电量",
                    float(a["E"][-1]),
                    None,
                ],
                [None] * 6,
            ]
        )
    table3 = [[item for d in DATES for item in (d + " 时段", "购电量 / kWh")]]
    for k in range(max(1, max(len(chosen[d][3]) for d in DATES))):
        cells = []
        for d in DATES:
            events = chosen[d][3]
            cells.extend(
                events[k]
                if k < len(events)
                else ["无", 0.0] if k == 0 else [None, None]
            )
        table3.append(cells)
    summary_rows = [
        [f"模型{model}：评价期汇总", None],
        ["评价期", "2025-02-01至2025-12-31（334天）"],
        [
            "模型定义",
            definition,
        ],
        ["比较口径", f"一月冻结方案：{MODEL_NAMES.get(selected, selected)}；本表方案：{model}"],
        ["指标", "数值"],
    ]
    summary_rows.extend([[label, summary[key]] for label, key in METRICS])
    if protocol.get("schema") == "statistical-risk-v1":
        risk_metrics = read_json(directory/"summary.json")
        for label, key in (("紧急费用CVaR / 元", "emergency_cvar"),
                           ("开发压力费用CVaR / 元", "development_stress_cvar"),
                           ("风险限额不满足天数", "risk_budget_unmet_days"),
                           ("最大物理残差 / kWh", "max_physical_residual")):
            summary_rows.append([label, risk_metrics[key]])
        summary_rows.extend([["CVaR水平", protocol["config"]["risk_beta"]],
                             ["一月选型费用增幅上限", protocol["config"]["cost_premium"]],
                             ["风险解释", "历史场景和人工压力检验；无限紧急补购，不代表停电可靠性"],
                             ["数据来源", "附件1.xlsx（电价）、附件2.xlsx（逐日负荷与光伏）"]])
    paper = {
        "评价期汇总": summary_rows,
        "表1_购电结果": table1,
        "表2_储能结果": table2,
        "表3_紧急购电": table3,
    }
    payload = dict(
        model=model,
        variant=variant,
        source=str(source),
        output=str(output),
        template=str(Path(protocol.get("attachments", ATTACHMENTS)) / "附件5/result2.xlsx"),
        summary=summary,
        purchase=purchase,
        storage=storage,
        emergency=emergency,
        paper=paper,
    )
    output.mkdir(parents=True, exist_ok=True)
    audit = output / "核验记录"
    audit.mkdir(exist_ok=True)
    write_json(audit / "tables.json", payload)
    write_json(
        audit / "数据来源.json",
        dict(
            model=model,
            variant=variant,
            source=str(source),
            day_count=len(daily),
            files={
                p.name: file_hash(p) for p in sorted(directory.glob("2025-??-??.npz"))
            },
        ),
    )
    (output / "使用说明.md").write_text(
        f"# 模型{model}提交材料\n\n"
        "- result2.xlsx：官方模板的全年计划购电、充放电和紧急购电三个工作表。\n"
        "- 论文结果表.xlsx：全年汇总及四个指定日期的表1、表2、表3。\n\n"
        f"方案：{variant}。总购电费 {summary['cash_cost']:,.4f} 元，紧急购电量 {summary['H']:,.4f} kWh，应急 {summary['emergency_days']} 天。\n\n"
        "计划购电量为G，全天总购电费含计划费用和五倍电价的紧急购电费用；所有电量单位为kWh。充放电量按每四小时汇总。连续应急十分钟时段合并，无事件日期记“无”。\n\n"
        "延续已核实的模板修正：144个时段从0:00-0:10到23:50-24:00。原模板不改动。Excel仅将显示精度设为四位小数，存储原始精度。\n\n"
        f"一月冻结方案为 {selected}；本次导出方案为 {variant}。模型定义与风险参数见计算目录中的 protocol_frozen.json。\n",
        encoding="utf-8",
    )
    return payload


def verify(source, output, payload):
    """重新读入导出的数值；openpyxl 只读，不用于生成工作簿。"""
    import openpyxl

    output = Path(output)
    count = 0
    for filename, expected in (
        (
            "result2.xlsx",
            dict(
                zip(
                    ("计划购电量", "充放电量", "紧急购电量"),
                    (payload["purchase"], payload["storage"], payload["emergency"]),
                )
            ),
        ),
        ("论文结果表.xlsx", payload["paper"]),
    ):
        wb = openpyxl.load_workbook(output / filename, data_only=True)
        if wb.sheetnames != list(expected):
            raise ValueError(f"工作表不完整：{filename}")
        for name, rows in expected.items():
            sheet = wb[name]
            for r, values in enumerate(rows, 1):
                for c, want in enumerate(values, 1):
                    cell = sheet.cell(r, c)
                    got = cell.value
                    if cell.data_type == "e":
                        raise ValueError(f"Excel错误：{name}!{cell.coordinate}")
                    if (
                        filename == "result2.xlsx"
                        and c == 1
                        and r > 1
                        and want is not None
                    ):
                        want = (date(1899, 12, 30) + timedelta(days=want)).isoformat()
                        got = got.date().isoformat()
                    if isinstance(want, (int, float)):
                        label = rows[r-1][0] if name == "评价期汇总" else ""
                        expected_format = ("0" if "天数" in str(label) or "时段数" in str(label)
                                           else "0.0000%" if label == "光伏消纳率" else "0.0000")
                        if cell.number_format != expected_format:
                            raise ValueError(f"显示精度错误：{name}!{cell.coordinate}")
                        if not isinstance(got, (int, float)) or abs(got - want) > 1e-7:
                            raise ValueError(f"数值不一致：{name}!{cell.coordinate}")
                    elif got != want:
                        raise ValueError(
                            f"内容不一致：{name}!{cell.coordinate}: {got!r}, {want!r}"
                        )
                    count += 1
        wb.close()
    variant, directory, rows = daily_results(source, payload["variant"])
    # Recheck source aggregate and full emergency coverage independently of presentation rounding.
    h_total = sum(
        float(read_arrays(directory / (r["date"] + ".npz"))["H"].sum()) for r in rows
    )
    exported = sum(r[2] for r in payload["emergency"][1:])
    if abs(h_total - exported) > 1e-5:
        raise ValueError("应急区间合计与全年结果不符")
    result = dict(
        model=payload["model"],
        days=len(rows),
        cells_checked=count,
        cash_cost=payload["summary"]["cash_cost"],
        emergency_kwh=h_total,
        verified=True,
        files={
            name: file_hash(output / name)
            for name in ("result2.xlsx", "论文结果表.xlsx")
        },
    )
    write_json(output / "核验记录/核验结果.json", result)
    return result


def export_excel(source, output, node=None, model=None, render=False):
    payload = make_payload(source, output, model)
    runtime = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
    node = Path(node or os.environ.get("CODEX_NODE", runtime / "node/bin/node.exe"))
    builder = CODE_DIR / "交表工具/export_submission.mjs"
    subprocess.run(
        [str(node), str(builder), str(Path(output).resolve() / "核验记录/tables.json")]
        + (["--render"] if render else []),
        check=True,
    )
    return verify(source, output, payload)


def export_with_runtime(source, output):
    """运行入口调用：表格工作使用独立的文档运行环境。"""
    runtime = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
    python = os.environ.get("CODEX_ARTIFACT_PYTHON", str(runtime / "python/python.exe"))
    subprocess.run(
        [
            python,
            "-X",
            "utf8",
            "-m",
            "microgrid.submission",
            "--source",
            str(Path(source).resolve()),
            "--output",
            str(Path(output).resolve()),
        ],
        cwd=CODE_DIR,
        check=True,
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    print(
        json.dumps(export_excel(args.source, args.output), ensure_ascii=False, indent=2)
    )

