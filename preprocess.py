from datetime import datetime, time
from pathlib import Path
import re

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook


DELTA_T = 1 / 6
EXPECTED_PERIODS = 144
BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "附件" / "附件1.xlsx"
OUTPUT_DIR = BASE_DIR / "outputs"


def clean_name(value):
    """简化列名，便于识别空格、换行和单位差异。"""
    text = "" if value is None else str(value)
    return re.sub(r"[\s_（）()\[\]【】/\\·]", "", text).lower()


def identify_field(value):
    name = clean_name(value)
    if "时间" in name or name in {"time", "时刻"}:
        return "time_raw"
    if "电价" in name or "price" in name:
        return "price"
    if "光伏" in name or name.startswith("pv"):
        return "pv_kw"
    if ("负载" in name or "负荷" in name or "load" in name) and "净" not in name:
        return "load_kw"
    return None


# 读取工作簿并自动寻找包含四个目标字段的表头行
def load_excel(path=INPUT_FILE):
    if not path.exists():
        raise FileNotFoundError(f"未找到输入文件：{path}")

    workbook = load_workbook(path, data_only=True, read_only=False)
    candidates = []
    for sheet in workbook.worksheets:
        for row in range(1, min(sheet.max_row, 30) + 1):
            fields = [identify_field(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)]
            score = len({field for field in fields if field})
            if score == 4:
                candidates.append((sheet.title, row))

    if not candidates:
        raise ValueError("未找到同时包含时间、电价、负载和光伏功率的表头行。")
    sheet_name, header_row = candidates[0]
    data = pd.read_excel(path, sheet_name=sheet_name, header=header_row - 1, engine="openpyxl")
    data = data.dropna(how="all").reset_index(drop=True)
    return data, workbook, sheet_name, header_row


# 打印 Excel 的实际结构，不修改原工作簿
def inspect_excel(workbook, sheet_name, header_row):
    sheet = workbook[sheet_name]
    blank_rows = [
        row for row in range(1, sheet.max_row + 1)
        if all(sheet.cell(row, col).value is None for col in range(1, sheet.max_column + 1))
    ]
    note_rows = []
    for row in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)]
        if any(value is not None for value in values) and sum(value is not None for value in values) <= 1:
            note_rows.append(row)

    headers = [sheet.cell(header_row, col).value for col in range(1, sheet.max_column + 1)]
    print("\n========== Excel 结构检查 ==========")
    print(f"工作表：{workbook.sheetnames}")
    print(f"采用工作表：{sheet_name}")
    print(f"有效区域：{sheet.calculate_dimension()}")
    print(f"表头行：第 {header_row} 行；数据起始行：第 {header_row + 1} 行")
    print(f"实际字段：{headers}")
    print(f"数据行数：{sheet.max_row - header_row}")
    print(f"合并单元格：{list(sheet.merged_cells.ranges) or '无'}")
    print(f"空白行：{blank_rows or '无'}")
    print(f"疑似备注行：{note_rows or '无'}")


# 按字段含义映射为统一列名
def standardize_columns(data):
    mapping = {}
    for column in data.columns:
        field = identify_field(column)
        if field:
            if field in mapping.values():
                raise ValueError(f"字段重复识别：{field}")
            mapping[column] = field

    required = {"time_raw", "price", "load_kw", "pv_kw"}
    missing = required - set(mapping.values())
    if missing:
        raise ValueError(f"缺少必要字段：{sorted(missing)}")
    return data[list(mapping)].rename(columns=mapping).copy(), mapping


def raw_time_text(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return "" if pd.isna(value) else str(value).strip()


def parse_time_minutes(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.hour * 60 + value.minute
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float, np.number)) and not pd.isna(value):
        return int(round(float(value) * 24 * 60)) if 0 <= float(value) <= 1 else np.nan

    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*(?:\+(\d+))?\s*", str(value))
    if not match:
        return np.nan
    hour, minute, days = int(match.group(1)), int(match.group(2)), int(match.group(3) or 0)
    return days * 1440 + hour * 60 + minute if hour < 24 and minute < 60 else np.nan


def minute_label(minutes):
    if minutes == 1440:
        return "24:00"
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


# 将附件中的区间结束时刻对齐为 144 个连续的 10 分钟区间
def build_time_intervals(data):
    raw_values = data["time_raw"].copy()
    parsed = raw_values.map(parse_time_minutes)
    periods = np.arange(1, len(data) + 1)
    expected_end = periods * 10

    data["time_raw"] = raw_values.map(raw_time_text)
    data.insert(0, "t", periods)
    data["time_start"] = [minute_label((t - 1) * 10) for t in periods]
    data["time_end"] = [minute_label(t * 10) for t in periods]
    data["time_interval"] = data["time_start"] + "-" + data["time_end"]

    valid = parsed.notna()
    conflicts = int((parsed[valid].to_numpy() != expected_end[valid.to_numpy()]).sum())
    parsed_values = parsed.dropna().astype(int)
    time_report = {
        "unparsed": int(parsed.isna().sum()),
        "conflicts": conflicts,
        "duplicates": int(parsed_values.duplicated().sum()),
        "continuous": bool(len(parsed_values) == len(data) and np.array_equal(parsed_values.to_numpy(), expected_end)),
        "covers_24h": bool(len(parsed_values) == EXPECTED_PERIODS and parsed_values.iloc[0] == 10 and parsed_values.iloc[-1] == 1440),
        "missing_ends": sorted(set(range(10, 1441, 10)) - set(parsed_values)),
    }
    return data, time_report


# 检查数据质量，只报告问题，不删除、插值或修正原始数值
def validate_data(data, time_report):
    invalid_numeric = {}
    for column in ["price", "load_kw", "pv_kw"]:
        original = data[column]
        converted = pd.to_numeric(original, errors="coerce")
        invalid_numeric[column] = int((original.notna() & converted.isna()).sum())
        data[column] = converted

    missing = data[["time_raw", "price", "load_kw", "pv_kw"]].isna().sum()
    duplicates = int(data[["time_raw", "price", "load_kw", "pv_kw"]].duplicated().sum())
    print("\n========== 数据质量检查 ==========")
    print(f"数据行数为 144：{'是' if len(data) == EXPECTED_PERIODS else '否'}（实际 {len(data)}）")
    print(f"缺失值：{missing.to_dict()}")
    print(f"完全重复数据行：{duplicates}")
    print(f"时间连续：{'是' if time_report['continuous'] else '否'}")
    print(f"恰好覆盖 24 小时：{'是' if time_report['covers_24h'] else '否'}")
    print(f"重复时间：{time_report['duplicates']}；缺失时段：{time_report['missing_ends'] or '无'}")
    print(f"时间无法解析：{time_report['unparsed']}；时间含义冲突：{time_report['conflicts']}")
    print(f"负电价：{int((data['price'] < 0).sum())}")
    print(f"负负载功率：{int((data['load_kw'] < 0).sum())}")
    print(f"负光伏功率：{int((data['pv_kw'] < 0).sum())}")
    print(f"无法解析的数值：{invalid_numeric}")
    print(f"数值字段类型正确：{all(pd.api.types.is_numeric_dtype(data[c]) for c in ['price', 'load_kw', 'pv_kw'])}")
    if time_report["conflicts"]:
        print("警告：附件时间与“区间结束时刻”的解释存在冲突，标准区间仍按行号生成，请人工核对。")
    return data


# 增加净负荷和仅用于统计的电量字段
def add_derived_features(data):
    data["net_load_kw"] = data["load_kw"] - data["pv_kw"]
    data["load_kwh"] = data["load_kw"] * DELTA_T
    data["pv_kwh"] = data["pv_kw"] * DELTA_T
    data["net_load_kwh"] = data["net_load_kw"] * DELTA_T
    columns = [
        "t", "time_raw", "time_start", "time_end", "time_interval",
        "price", "load_kw", "pv_kw", "net_load_kw",
        "load_kwh", "pv_kwh", "net_load_kwh",
    ]
    return data[columns]


def peak_text(data, column, kind="max"):
    series = data[column].dropna()
    if series.empty:
        return "无有效数据"
    index = series.idxmax() if kind == "max" else series.idxmin()
    return f"{data.loc[index, 'time_interval']}（{data.loc[index, column]:.4f}）"


# 输出问题一建模前需要的基础统计
def calculate_statistics(data):
    load_energy = data["load_kwh"].sum(min_count=1)
    pv_energy = data["pv_kwh"].sum(min_count=1)
    surplus = data.loc[data["pv_kw"] > data["load_kw"], "time_interval"].tolist()
    ratio = load_energy / pv_energy if pd.notna(pv_energy) and pv_energy != 0 else np.nan

    print("\n========== 基础数据摘要 ==========")
    print(f"数据条数：{len(data)}")
    print(f"时间范围：{data['time_interval'].iloc[0]} 至 {data['time_interval'].iloc[-1]}")
    print(f"电价（元/kWh）：最小 {data['price'].min():.4f}，最大 {data['price'].max():.4f}，平均 {data['price'].mean():.4f}")
    print(f"最低电价时段：{peak_text(data, 'price', 'min')}")
    print(f"最高电价时段：{peak_text(data, 'price', 'max')}")
    print(f"负载（kW）：最小 {data['load_kw'].min():.4f}，最大 {data['load_kw'].max():.4f}，平均 {data['load_kw'].mean():.4f}")
    print(f"负载峰值时段：{peak_text(data, 'load_kw')}；全天负荷电量：{load_energy:.4f} kWh")
    print(f"光伏（kW）：最大 {data['pv_kw'].max():.4f}，平均 {data['pv_kw'].mean():.4f}")
    print(f"光伏峰值时段：{peak_text(data, 'pv_kw')}；全天预测发电量：{pv_energy:.4f} kWh")
    print(f"净负荷（kW）：最小 {data['net_load_kw'].min():.4f}，最大 {data['net_load_kw'].max():.4f}，平均 {data['net_load_kw'].mean():.4f}")
    print(f"净负荷最大时段：{peak_text(data, 'net_load_kw')}")
    print(f"净负荷小于 0 的时段数：{int((data['net_load_kw'] < 0).sum())}")
    print(f"光伏大于负载的时段：{', '.join(surplus) if surplus else '无'}")
    print(f"全天负荷总电量 / 全天光伏预测发电量：{ratio:.4f}" if pd.notna(ratio) else "该比值无法计算")


def set_plot_font():
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def draw_curve(data, columns, labels, title, ylabel, filename):
    x = data["t"].to_numpy()
    fig, ax = plt.subplots(figsize=(10, 5.5))
    for column, label in zip(columns, labels):
        ax.plot(x, data[column], label=label, linewidth=1.5)
    ticks = list(range(1, len(data) + 1, 12))
    ax.set_xticks(ticks, data.loc[[tick - 1 for tick in ticks], "time_start"])
    ax.set_xlabel("时间")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.grid(alpha=0.3)
    if len(columns) > 1:
        ax.legend()
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / filename, dpi=300, bbox_inches="tight")
    plt.close(fig)


# 分别保存论文可用的基础曲线图
def plot_data(data):
    set_plot_font()
    draw_curve(data, ["price"], ["电价"], "分时电价曲线", "电价（元/kWh）", "price_curve.png")
    draw_curve(data, ["load_kw", "pv_kw"], ["小区负载", "光伏预测功率"], "负载与光伏预测功率", "功率（kW）", "load_pv_curve.png")
    draw_curve(data, ["net_load_kw"], ["净负荷"], "净负荷曲线", "净负荷（kW）", "net_load_curve.png")
    draw_curve(data, ["load_kw", "pv_kw", "net_load_kw"], ["小区负载", "光伏预测功率", "净负荷"], "基础功率对比", "功率（kW）", "basic_energy_curve.png")


# 保存 CSV 和 Excel，保留功率字段并附加电量统计字段
def save_processed_data(data):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUTPUT_DIR / "preprocessed_q1.csv"
    xlsx_path = OUTPUT_DIR / "preprocessed_q1.xlsx"
    data.to_csv(csv_path, index=False, encoding="utf-8-sig", float_format="%.6f")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="preprocessed_q1")
        sheet = writer.book["preprocessed_q1"]
        sheet.freeze_panes = "A2"
        widths = [8, 12, 12, 12, 18, 14, 14, 14, 16, 14, 14, 18]
        for column, width in zip(sheet.columns, widths):
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2, min_col=6, max_col=12):
            for cell in row:
                cell.number_format = "0.000000"
    return csv_path, xlsx_path


def main():
    data, workbook, sheet_name, header_row = load_excel()
    inspect_excel(workbook, sheet_name, header_row)
    data, mapping = standardize_columns(data)
    print(f"字段映射：{mapping}")
    data, time_report = build_time_intervals(data)
    data = validate_data(data, time_report)
    data = add_derived_features(data)
    calculate_statistics(data)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    plot_data(data)
    csv_path, xlsx_path = save_processed_data(data)
    print("\n========== 输出文件 ==========")
    print(csv_path)
    print(xlsx_path)
    print(f"图像已保存至：{OUTPUT_DIR}")


if __name__ == "__main__":
    main()
