"""Re-execute locked plans against raw observations, then export auditable metrics.

This is an execution/accounting replay, NOT a fresh day-ahead optimization.
The saved H/C/D/U/W/fees are used only for independent comparison, never as inputs.
"""
from pathlib import Path
from datetime import datetime
import argparse,hashlib,json
import numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Border,Font,Side
from openpyxl.utils import get_column_letter
from q2_config import HERE,load_config
from q2_data import load_data
from q2_controller import rollout
from q2_backtest import cold_start
from q2_archive import save_json,save_npz

MODELS=('M0','M1','M2','M3')
METRICS=[
 ('计划购电量','kWh','G'),('紧急购电量','kWh','H'),
 ('总购电量（计划＋紧急）','kWh','total_purchase'),
 ('计划购电费','元','plan_fee'),('紧急购电费','元','emergency_fee'),
 ('实际总购电费','元','cash_fee'),('发生应急的天数','日','emergency_days'),
 ('应急时段数','个','emergency_slots'),('未利用计划额度','kWh','U'),
 ('弃光量','kWh','W'),('储能转换损耗','kWh','storage_loss'),
 ('光伏消纳率','%','pv_utilization'),('12月31日24:00储电量','kWh','Eend')]

def daily_values(date,model,o,price,cfg):
    C,D=float(o['charge'].sum()),float(o['discharge'].sum())
    # Microgrid-side C,D: input minus useful output minus inventory increase.
    loss=(1-cfg.eta_c)*C+(1/cfg.eta_d-1)*D
    conservation_loss=C-D-float(o['E'][-1]-o['E'][0])
    assert abs(loss-conservation_loss)<1e-7
    H=float(o['H'].sum())
    plan_fee=float(price@o['G'])
    emergency_fee=float(cfg.emergency_multiplier*price@o['H'])
    return dict(date=date,model=model,G=float(o['G'].sum()),H=H,U=float(o['U'].sum()),
        W=float(o['W'].sum()),load=float(o['load'].sum()),pv=float(o['pv'].sum()),
        charge=C,discharge=D,E0=float(o['E'][0]),Eend=float(o['E'][-1]),
        storage_loss=loss,plan_fee=plan_fee,emergency_fee=emergency_fee,
        cash_fee=plan_fee+emergency_fee,
        emergency_days=int(np.any(o['H']>cfg.physical_tolerance)),
        emergency_slots=int(np.count_nonzero(o['H']>cfg.physical_tolerance)))

def aggregate(rows):
    keys=('G','H','U','W','load','pv','charge','discharge','storage_loss',
          'plan_fee','emergency_fee','cash_fee','emergency_days','emergency_slots')
    result={k:sum(r[k] for r in rows) for k in keys}
    result.update(days=len(rows),E0=rows[0]['E0'],Eend=rows[-1]['Eend'])
    result['total_purchase']=result['G']+result['H']
    result['pv_utilization']=1-result['W']/result['pv']
    result['balance_error']=abs(result['G']-result['U']+result['pv']-result['W']+result['H']-
        result['load']-(result['Eend']-result['E0'])-result['storage_loss'])
    assert result['balance_error']<1e-6
    assert abs(result['cash_fee']-result['plan_fee']-result['emergency_fee'])<1e-6
    return result

def export_workbook(path,formal,annual,daily,run_time):
    book=Workbook()
    book.remove(book.active)
    for name,period,values in [('正式334天对比','2025-02-01至2025-12-31，不含1月预热',formal),
                               ('全年365天对比','2025-01-01至2025-12-31，包含共同冷启动和预热',annual)]:
        sheet=book.create_sheet(name)
        sheet.append(['四模型评价期结果'])
        sheet.merge_cells('A1:F1')
        sheet.append([period]);sheet.merge_cells('A2:F2')
        sheet.append(['指标','单位',*MODELS])
        for label,unit,key in METRICS:
            sheet.append([label,unit,*[values[m][key] for m in MODELS]])
        sheet.column_dimensions['A'].width=34
        sheet.column_dimensions['B'].width=9
        for column in 'CDEF':sheet.column_dimensions[column].width=23
        sheet.row_dimensions[1].height=30
        sheet.row_dimensions[2].height=25
        for row in sheet.iter_rows(min_row=3):
            sheet.row_dimensions[row[0].row].height=29
        line=Side(style='thin',color='222222')
        for row in sheet:
            for cell in row:
                cell.font=Font(name='Microsoft YaHei',size=11,bold=cell.row in (1,3))
                cell.alignment=Alignment(vertical='center',horizontal='left' if cell.column<3 else 'right')
                if cell.row==3:cell.border=Border(top=line,bottom=line)
                if cell.row==sheet.max_row:cell.border=Border(bottom=line)
        for i,(_,unit,key) in enumerate(METRICS,4):
            for column in range(3,7):
                sheet.cell(i,column).number_format='0.00%' if unit=='%' else ('0' if unit in ('日','个') else '#,##0.00')
        sheet.sheet_view.showGridLines=False
        sheet.print_area='A1:F16'
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
    detail=book.create_sheet('逐日计算明细')
    fields=[('model','模型'),('date','日期'),('G','计划购电量/kWh'),('H','紧急购电量/kWh'),
        ('plan_fee','计划购电费/元'),('emergency_fee','紧急购电费/元'),('cash_fee','实际总购电费/元'),
        ('emergency_days','是否发生应急'),('emergency_slots','应急时段数'),('U','未利用计划额度/kWh'),
        ('W','弃光量/kWh'),('storage_loss','储能转换损耗/kWh'),('pv','光伏发电量/kWh'),
        ('charge','充电量/kWh'),('discharge','放电量/kWh'),('E0','0:00储电量/kWh'),('Eend','24:00储电量/kWh')]
    detail.append([label for _,label in fields])
    for row in daily:
        detail.append([datetime.fromisoformat(row[k]) if k=='date' else row[k] for k,_ in fields])
    detail.freeze_panes='C2';detail.auto_filter.ref=detail.dimensions
    detail.sheet_view.showGridLines=False
    for i in range(1,len(fields)+1):detail.column_dimensions[get_column_letter(i)].width=23
    detail.column_dimensions['A'].width=15;detail.column_dimensions['B'].width=15
    for row in detail:
        for cell in row:
            cell.font=Font(name='Microsoft YaHei',size=10,bold=cell.row==1)
            cell.alignment=Alignment(vertical='center',wrap_text=cell.row==1)
            if cell.row>1:
                cell.number_format='yyyy-mm-dd' if cell.column==2 else ('0' if cell.column in (8,9) else '#,##0.00')
    detail.row_dimensions[1].height=35
    notes=book.create_sheet('计算口径')
    for row in [
        ['项目','说明'],
        ['本次计算方式','使用已锁定计划和参考轨迹，重新读取实际供需逐槽执行控制器；没有重新优化购电计划。'],
        ['计算时间',run_time],
        ['正式评价期','2025-02-01至2025-12-31，334天，每天144槽。'],
        ['365天口径','1月1日零计划、电池不动作；1月2—31日共同预热，各模型全年结果均加入同一段预热。'],
        ['计划购电费','逐槽电价 × 锁定计划G，全年求和；未使用额度仍付费。'],
        ['紧急购电费','逐槽电价 × 紧急电量H × 5，全年求和。'],
        ['实际总购电费','计划购电费＋紧急购电费；不包含终端价值项。'],
        ['总购电量','计划G＋紧急H，包含已付费而未使用的计划额度。'],
        ['应急天数和时段数','按H > 0.000001 kWh计数；费用与电量汇总保留全部H，不截去微量电量。'],
        ['储能转换损耗','(1-0.9) × 总充电量 ＋ (1/0.9-1) × 总放电量；充放电量均为微网侧口径。'],
        ['损耗独立校验','总充电量－总放电量－(期末内部储电量－期初内部储电量)，应与转换损耗一致。'],
        ['光伏消纳率','1－弃光总量/光伏发电总量；不对每日消纳率直接平均。'],
        ['年末储电量','实际执行轨迹在2025-12-31 24:00的内部电量，单位kWh。'],
        ['计算输入','附件1电价、附件2实际负载和光伏、四模型已锁定计划与参考轨迹。'],
        ['截图用途','仅采用指标和表格展示形式，截图旧版数值没有用于本次计算。'],
    ]:notes.append(row)
    notes.column_dimensions['A'].width=25;notes.column_dimensions['B'].width=92
    notes.sheet_view.showGridLines=False
    for row in notes:
        notes.row_dimensions[row[0].row].height=40 if row[0].row>1 else 28
        for cell in row:
            cell.font=Font(name='Microsoft YaHei',size=10,bold=cell.row==1)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
    path.parent.mkdir(parents=True,exist_ok=True)
    book.save(path)
    # Re-read every headline cell and daily inputs, not merely the displayed strings.
    checked=load_workbook(path,data_only=True)
    for name,values in [('正式334天对比',formal),('全年365天对比',annual)]:
        for r,(_,_,key) in enumerate(METRICS,4):
            for c,m in enumerate(MODELS,3):
                assert abs(checked[name].cell(r,c).value-values[m][key])<1e-6
    assert checked['逐日计算明细'].max_row==len(daily)+1
    for row,expected in zip(checked['逐日计算明细'].iter_rows(min_row=2,values_only=True),daily,strict=True):
        for value,(key,_) in zip(row,fields):
            if key=='date': assert value.date().isoformat()==expected[key]
            elif isinstance(expected[key],str):assert value==expected[key]
            else:assert abs(value-expected[key])<1e-6

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--run',default=str(HERE/'result/formal_v2'))
    parser.add_argument('--output',default=str(HERE/'result/recalculated_metrics'))
    parser.add_argument('--xlsx',default=str(HERE/'result/四模型指标重算.xlsx'))
    args=parser.parse_args()
    root,out=Path(args.run),Path(args.output)
    cfg=load_config(root/'config.json')
    data=load_data()
    run_time=datetime.now().astimezone().isoformat(timespec='seconds')
    if out.exists():raise FileExistsError('Choose a new calculation directory; existing evidence is immutable.')
    save_json(out/'execution.json',dict(mode='locked_plan_causal_replay_and_independent_accounting',
        fresh_day_ahead_optimization=False,started_at=run_time,config_hash=cfg.version,
        models=list(MODELS),formal_days=334,slots_per_day=144,emergency_count_threshold_kwh=cfg.physical_tolerance))
    source_hashes={}
    def remember(path):source_hashes[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest()
    for path in [root/'config.json',HERE.parent/'附件/附件1.xlsx',HERE.parent/'附件/附件2.xlsx']:remember(path)
    maximum_error=0.;maximum_constraint=0.
    def execute_day(date,model,E0):
        nonlocal maximum_error,maximum_constraint
        strategy='warmup' if model=='共同预热' else 'star_'+model
        planpath=root/'locked_plans'/strategy/(date+'.npz')
        remember(planpath)
        with np.load(planpath) as p:
            G,ref=p['G'].copy(),p['reference'].copy()
        original_G=G.copy()
        o=rollout(G,ref,E0,iter(data.values[data.dates.index(date)]),data.price,cfg)
        np.testing.assert_array_equal(G,original_G)
        recordpath=root/'runs'/strategy/(date+'.npz')
        remember(recordpath)
        with np.load(recordpath) as original:
            for key in ['G','H','charge','discharge','U','W','E','load','pv']:
                error=float(np.abs(original[key]-o[key]).max())
                assert error<1e-7,(date,model,key,error)
                maximum_error=max(maximum_error,error)
        maximum_constraint=max(maximum_constraint,max(o['residuals'].values()))
        save_npz(out/'replayed_slots'/strategy/(date+'.npz'),**{k:v for k,v in o.items() if isinstance(v,np.ndarray)})
        return o,daily_values(date,model,o,data.price,cfg)
    cold=cold_start(data.values[0],data.price,cfg)
    warm=[daily_values(data.dates[0],'共同预热',cold,data.price,cfg)]
    save_npz(out/'replayed_slots/warmup/2025-01-01.npz',**{k:v for k,v in cold.items() if isinstance(v,np.ndarray)})
    energy=cfg.initial_energy
    for date in data.dates[1:31]:
        o,row=execute_day(date,'共同预热',energy)
        energy=float(o['E'][-1]);warm.append(row)
    frozen=json.loads((root/'freeze.json').read_text(encoding='utf-8'))
    assert abs(energy-frozen['shared_E0'])<1e-7
    formal,annual,all_daily={},{},warm.copy()
    for model in MODELS:
        state=energy;rows=[]
        for date in data.dates[31:]:
            o,row=execute_day(date,model,state)
            state=float(o['E'][-1]);rows.append(row)
        assert len(rows)==334
        formal[model]=aggregate(rows)
        annual[model]=aggregate(warm+rows)
        all_daily+=rows
        print(json.dumps(dict(model=model,days=len(rows),slots=len(rows)*144,
            **{k:formal[model][k] for k in ['G','H','cash_fee','storage_loss','pv_utilization','Eend']}),ensure_ascii=False),flush=True)
    for path,expected in source_hashes.items():assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==expected,path
    save_json(out/'source_hashes.json',source_hashes)
    save_json(out/'recalculated_metrics.json',dict(formal334=formal,calendar365=annual,warmup31=aggregate(warm)))
    save_json(out/'daily_metrics.json',all_daily)
    export_workbook(Path(args.xlsx),formal,annual,all_daily,run_time)
    save_json(out/'verification.json',dict(completed=True,finished_at=datetime.now().astimezone().isoformat(timespec='seconds'),
        model_days=1336,formal_slots=192384,warmup_days=31,
        source_files_unchanged=True,max_replay_difference=maximum_error,
        max_physical_residual=maximum_constraint,spreadsheet_reread_passed=True,
        compared_original_summaries=False,all_metrics_computed_from_replayed_slots=True,
        workbook_sha256=hashlib.sha256(Path(args.xlsx).read_bytes()).hexdigest()))
    print('Completed workbook: '+args.xlsx,flush=True)

if __name__=='__main__':main()
